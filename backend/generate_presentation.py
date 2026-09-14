import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_nexus_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    # Colors
    BG_LIGHT = RGBColor(248, 250, 252)       # #F8FAFC
    WHITE = RGBColor(255, 255, 255)          # #FFFFFF
    TEXT_DARK = RGBColor(15, 23, 42)         # #0F172A (Slate 900)
    TEXT_MUTED = RGBColor(71, 85, 105)       # #475569 (Slate 600)
    ACCENT_EMERALD = RGBColor(16, 185, 129)  # #10B981
    ACCENT_INDIGO = RGBColor(79, 70, 229)    # #4F46E5
    ACCENT_AMBER = RGBColor(245, 158, 11)    # #F59E0B
    BORDER_COLOR = RGBColor(226, 232, 240)   # #E2E8F0
    CARD_BG = RGBColor(255, 255, 255)
    CARD_SUBTLE_BG = RGBColor(241, 245, 249) # #F1F5F9

    def set_slide_background(slide, color=BG_LIGHT):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, category, title, subtitle=None):
        # Category Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.8), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor(236, 253, 245)
        badge.line.color.rgb = RGBColor(167, 243, 208)
        tf_badge = badge.text_frame
        tf_badge.word_wrap = True
        p_badge = tf_badge.paragraphs[0]
        p_badge.text = category.upper()
        p_badge.font.size = Pt(10)
        p_badge.font.bold = True
        p_badge.font.color.rgb = RGBColor(5, 150, 105)
        p_badge.alignment = PP_ALIGN.CENTER

        # Main Title
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK

        if subtitle:
            tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.4))
            tf_sub = tb_sub.text_frame
            tf_sub.word_wrap = True
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.size = Pt(13)
            p_sub.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, BG_LIGHT)

    # Hero Banner Box
    add_card(s1, 0.8, 1.2, 11.7, 5.0, WHITE, BORDER_COLOR)

    # Pill Tag
    tag = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.7), Inches(3.4), Inches(0.4))
    tag.fill.solid()
    tag.fill.fore_color.rgb = RGBColor(238, 242, 255)
    tag.line.color.rgb = RGBColor(199, 210, 254)
    tf_tag = tag.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "OFFICIAL EXECUTIVE PRESENTATION"
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_INDIGO
    p_tag.alignment = PP_ALIGN.CENTER

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(10.9), Inches(1.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "NEXUS OS — RESOURCE GUARDIAN"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_DARK

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Autonomous AI System Telemetry & Self-Healing Operating System Agent"
    p1_sub.font.size = Pt(18)
    p1_sub.font.color.rgb = ACCENT_EMERALD
    p1_sub.space_before = Pt(10)

    # 3 Mini Feature Badges on Title Slide
    features = [
        ("⚡ Real-Time Telemetry", "2-second sampling interval for CPU, RAM & Disk"),
        ("🧠 ML IsolationForest", "Unsupervised anomaly detection & predictive forecasting"),
        ("🛡 Safe Guardrails", "Protected OS blacklist & instant SOC2 PDF export")
    ]
    for i, (feat_title, feat_desc) in enumerate(features):
        c_left = 1.2 + i * 3.7
        add_card(s1, c_left, 4.3, 3.4, 1.4, CARD_SUBTLE_BG, BORDER_COLOR)
        tb_f = s1.shapes.add_textbox(Inches(c_left + 0.1), Inches(4.4), Inches(3.2), Inches(1.2))
        tf_f = tb_f.text_frame
        tf_f.word_wrap = True
        pf1 = tf_f.paragraphs[0]
        pf1.text = feat_title
        pf1.font.size = Pt(13)
        pf1.font.bold = True
        pf1.font.color.rgb = TEXT_DARK
        pf2 = tf_f.add_paragraph()
        pf2.text = feat_desc
        pf2.font.size = Pt(11)
        pf2.font.color.rgb = TEXT_MUTED
        pf2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: Problem Statement
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "EXECUTIVE PROBLEM", "The Challenge of Uncontrolled System Exhaustion", "Modern computer systems and servers suffer from silent performance degradation and abrupt crashes.")

    problems = [
        ("💥 Uncontrolled Resource Spikes", "Software memory leaks or runaway background loops quickly consume 100% CPU or available RAM, leading to total OS unresponsiveness."),
        ("⏳ Slow Manual Intervention", "Human administrators and users notice degradation only after system freeze occurs, when opening Task Manager is impossible."),
        ("🛑 Risk of Collateral Damage", "Manual process termination risks accidentally killing core system dependencies, triggering kernel panics or data corruption.")
    ]
    for i, (ptitle, pdesc) in enumerate(problems):
        c_left = 0.8 + i * 3.9
        add_card(s2, c_left, 2.2, 3.7, 4.5, WHITE, BORDER_COLOR)
        tb_p = s2.shapes.add_textbox(Inches(c_left + 0.2), Inches(2.4), Inches(3.3), Inches(4.1))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True
        pp1 = tf_p.paragraphs[0]
        pp1.text = ptitle
        pp1.font.size = Pt(16)
        pp1.font.bold = True
        pp1.font.color.rgb = TEXT_DARK
        pp2 = tf_p.add_paragraph()
        pp2.text = pdesc
        pp2.font.size = Pt(13)
        pp2.font.color.rgb = TEXT_MUTED
        pp2.space_before = Pt(12)

    # =========================================================================
    # SLIDE 3: The Solution (Nexus OS)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "THE SOLUTION", "Nexus OS — Autonomous Health & Self-Healing Agent", "Proactive, deterministic, and AI-assisted system protection that acts before failure occurs.")

    sol_cards = [
        ("1. Proactive Telemetry", "Continuous psutil monitoring tracks per-process CPU, memory allocation, and IO operations every 2s."),
        ("2. Machine Learning AI", "IsolationForest algorithm evaluates anomaly contamination scores dynamically without pre-labeled data."),
        ("3. Immune Guardrails", "Protected OS process blacklist guarantees critical drivers & system files are never terminated."),
        ("4. Executive Transparency", "Live WebSockets dashboard paired with 1-click audit PDF reports for compliance verification.")
    ]
    for i, (stitle, sdesc) in enumerate(sol_cards):
        row = i // 2
        col = i % 2
        c_left = 0.8 + col * 5.9
        c_top = 2.2 + row * 2.4
        add_card(s3, c_left, c_top, 5.7, 2.1, WHITE, BORDER_COLOR)
        tb_s = s3.shapes.add_textbox(Inches(c_left + 0.2), Inches(c_top + 0.2), Inches(5.3), Inches(1.7))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        ps1 = tf_s.paragraphs[0]
        ps1.text = stitle
        ps1.font.size = Pt(16)
        ps1.font.bold = True
        ps1.font.color.rgb = ACCENT_INDIGO
        ps2 = tf_s.add_paragraph()
        ps2.text = sdesc
        ps2.font.size = Pt(13)
        ps2.font.color.rgb = TEXT_MUTED
        ps2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 4: System Architecture Diagram
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "SYSTEM ARCHITECTURE", "High-Level System Component & Dataflow Architecture", "Decoupled architecture connecting low-level OS telemetry with AI analytics and UI presentation.")

    # Architecture Nodes (Visual Diagram Boxes)
    nodes = [
        ("💻 Hardware & OS Layer", "psutil Telemetry Collector\nWindows / Linux OS APIs", 0.8, 2.2, 3.5, 2.0, CARD_SUBTLE_BG),
        ("⚡ FastAPI Core Service", "Async REST Controllers\nWebSocket Metric Streamer", 4.9, 2.2, 3.5, 2.0, RGBColor(238, 242, 255)),
        ("🧠 ML Analytics Engine", "Isolation Forest Model\nLinear / LSTM Forecaster", 9.0, 2.2, 3.5, 2.0, RGBColor(236, 253, 245)),
        ("🛡 Safety Guardrail Engine", "Blacklist / Whitelist Verifier\nSimulation Mode Manager", 4.9, 4.7, 3.5, 2.0, RGBColor(254, 242, 242)),
        ("📊 React 19 Frontend Console", "Live Dashboard & Process Table\n1-Click Executive PDF Exporter", 9.0, 4.7, 3.5, 2.0, WHITE)
    ]
    for title, desc, left, top, width, height, bg in nodes:
        add_card(s4, left, top, width, height, bg, BORDER_COLOR)
        tb_n = s4.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.2), Inches(width - 0.3), Inches(height - 0.4))
        tf_n = tb_n.text_frame
        tf_n.word_wrap = True
        pn1 = tf_n.paragraphs[0]
        pn1.text = title
        pn1.font.size = Pt(14)
        pn1.font.bold = True
        pn1.font.color.rgb = TEXT_DARK
        pn2 = tf_n.add_paragraph()
        pn2.text = desc
        pn2.font.size = Pt(11)
        pn2.font.color.rgb = TEXT_MUTED
        pn2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 5: The 4-Step Self-Healing Pipeline
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "AUTOMATION PIPELINE", "The 4-Step Self-Healing Pipeline", "Deterministic workflow ensuring complete safety before taking any automated system action.")

    steps = [
        ("STEP 1: INGEST", "Sampling", "Real-time process tree, CPU %, RAM pages & disk IO polled every 2 seconds."),
        ("STEP 2: DETECT", "ML Scoring", "IsolationForest calculates outlier contamination score for every active PID."),
        ("STEP 3: VERIFY", "Guardrails", "Checks blacklist protection, simulation mode state, and maintenance schedule."),
        ("STEP 4: HEAL", "Remediation", "Executes non-destructive Renice / Throttle or graceful process restart.")
    ]
    for i, (step_num, step_name, step_desc) in enumerate(steps):
        c_left = 0.8 + i * 2.95
        add_card(s5, c_left, 2.4, 2.7, 4.3, WHITE, BORDER_COLOR)
        
        # Step Header Pill
        spill = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left + 0.2), Inches(2.6), Inches(2.3), Inches(0.35))
        spill.fill.solid()
        spill.fill.fore_color.rgb = RGBColor(236, 253, 245)
        spill.line.color.rgb = RGBColor(167, 243, 208)
        tf_sp = spill.text_frame
        psp = tf_sp.paragraphs[0]
        psp.text = step_num
        psp.font.size = Pt(10)
        psp.font.bold = True
        psp.font.color.rgb = ACCENT_EMERALD
        psp.alignment = PP_ALIGN.CENTER

        tb_step = s5.shapes.add_textbox(Inches(c_left + 0.2), Inches(3.1), Inches(2.3), Inches(3.3))
        tf_step = tb_step.text_frame
        tf_step.word_wrap = True
        pt1 = tf_step.paragraphs[0]
        pt1.text = step_name
        pt1.font.size = Pt(16)
        pt1.font.bold = True
        pt1.font.color.rgb = TEXT_DARK
        pt2 = tf_step.add_paragraph()
        pt2.text = step_desc
        pt2.font.size = Pt(12)
        pt2.font.color.rgb = TEXT_MUTED
        pt2.space_before = Pt(10)

    # =========================================================================
    # SLIDE 6: Machine Learning Engine
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "MACHINE LEARNING", "Unsupervised Anomaly Detection & Forecasting", "Isolation Forest ML engine identifies complex multi-variable system anomalies dynamically.")

    ml_cards = [
        ("🌲 Isolation Forest Model", "Unsupervised machine learning that isolates anomalies instead of profiling normal points. Requires zero pre-labeled training data."),
        ("📈 10-Minute Predictive Horizon", "Linear regression and time-series forecasting predict memory exhaustion up to 10 minutes before system crash."),
        ("🔄 Dynamic Model Auto-Tuning", "Operator feedback on false positives dynamically recalibrates model contamination sensitivity parameters.")
    ]
    for i, (mtitle, mdesc) in enumerate(ml_cards):
        c_left = 0.8 + i * 3.9
        add_card(s6, c_left, 2.2, 3.7, 4.5, WHITE, BORDER_COLOR)
        tb_m = s6.shapes.add_textbox(Inches(c_left + 0.2), Inches(2.5), Inches(3.3), Inches(4.0))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        pm1 = tf_m.paragraphs[0]
        pm1.text = mtitle
        pm1.font.size = Pt(16)
        pm1.font.bold = True
        pm1.font.color.rgb = TEXT_DARK
        pm2 = tf_m.add_paragraph()
        pm2.text = mdesc
        pm2.font.size = Pt(13)
        pm2.font.color.rgb = TEXT_MUTED
        pm2.space_before = Pt(12)

    # =========================================================================
    # SLIDE 7: Safety Policies & Guardrails
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "SAFETY GUARDRAILS", "Deterministic Operating Policies & Immunity Lists", "Multi-layered safety boundaries guarantee that system-critical operations remain untouched.")

    g_items = [
        ("🛡 Protected System Blacklist", "Hardened immunity list covering systemd, explorer.exe, python, and uvicorn. Immune to auto-termination."),
        ("🧪 Simulation Mode Switch", "Safe dry-run operating mode where AI logs all healing decisions without making real system modifications."),
        ("⏰ Maintenance Window Schedules", "Time-based exemption policies (e.g. 02:00–04:00 AM) permitting heavy batch operations during off-peak hours.")
    ]
    for i, (gtitle, gdesc) in enumerate(g_items):
        c_left = 0.8 + i * 3.9
        add_card(s7, c_left, 2.2, 3.7, 4.5, WHITE, BORDER_COLOR)
        tb_g = s7.shapes.add_textbox(Inches(c_left + 0.2), Inches(2.5), Inches(3.3), Inches(4.0))
        tf_g = tb_g.text_frame
        tf_g.word_wrap = True
        pg1 = tf_g.paragraphs[0]
        pg1.text = gtitle
        pg1.font.size = Pt(16)
        pg1.font.bold = True
        pg1.font.color.rgb = TEXT_DARK
        pg2 = tf_g.add_paragraph()
        pg2.text = gdesc
        pg2.font.size = Pt(13)
        pg2.font.color.rgb = TEXT_MUTED
        pg2.space_before = Pt(12)

    # =========================================================================
    # SLIDE 8: Technology Stack
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "TECHNOLOGY STACK", "Modern Decoupled Software Stack", "Built with industry-standard, high-performance web and data science frameworks.")

    techs = [
        ("Frontend UI", "React 19, TypeScript, Vite, TailwindCSS, Recharts, jsPDF + html2canvas"),
        ("Backend Services", "Python 3.11+, FastAPI (Async), WebSockets, SQLAlchemy, Uvicorn"),
        ("Machine Learning", "Scikit-Learn (IsolationForest), Statsmodels, NumPy"),
        ("OS Integration", "psutil, Windows Win32 API / Linux POSIX process controllers")
    ]
    for i, (ttitle, tdesc) in enumerate(techs):
        row = i // 2
        col = i % 2
        c_left = 0.8 + col * 5.9
        c_top = 2.2 + row * 2.4
        add_card(s8, c_left, c_top, 5.7, 2.1, WHITE, BORDER_COLOR)
        tb_t = s8.shapes.add_textbox(Inches(c_left + 0.2), Inches(c_top + 0.2), Inches(5.3), Inches(1.7))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        pt1 = tf_t.paragraphs[0]
        pt1.text = ttitle
        pt1.font.size = Pt(16)
        pt1.font.bold = True
        pt1.font.color.rgb = ACCENT_INDIGO
        pt2 = tf_t.add_paragraph()
        pt2.text = tdesc
        pt2.font.size = Pt(13)
        pt2.font.color.rgb = TEXT_MUTED
        pt2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 9: Reporting & Compliance
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "AUDIT & REPORTING", "Executive PDF & SOC2 Compliance Reporting", "Comprehensive diagnostic reports with cryptographic verification signatures.")

    r_cards = [
        ("📊 Live Web Console", "Interactive per-process telemetry tables, slide-over drawer, and real-time metric stream."),
        ("📄 1-Click Executive PDF Export", "Formatted multi-page PDF documents containing vitals, guardrail states, and anomaly audit logs."),
        ("🔒 SHA-256 Digest Signature", "Cryptographic signature embedded on every generated report for audit compliance verifiability.")
    ]
    for i, (rtitle, rdesc) in enumerate(r_cards):
        c_left = 0.8 + i * 3.9
        add_card(s9, c_left, 2.2, 3.7, 4.5, WHITE, BORDER_COLOR)
        tb_r = s9.shapes.add_textbox(Inches(c_left + 0.2), Inches(2.5), Inches(3.3), Inches(4.0))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True
        pr1 = tf_r.paragraphs[0]
        pr1.text = rtitle
        pr1.font.size = Pt(16)
        pr1.font.bold = True
        pr1.font.color.rgb = TEXT_DARK
        pr2 = tf_r.add_paragraph()
        pr2.text = rdesc
        pr2.font.size = Pt(13)
        pr2.font.color.rgb = TEXT_MUTED
        pr2.space_before = Pt(12)

    # =========================================================================
    # SLIDE 10: Conclusion & Impact Comparison Table
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "CONCLUSION & IMPACT", "Business Impact & Value Comparison", "Transforming system maintenance from reactive emergency firefighting to autonomous stability.")

    # Table comparing Traditional vs Nexus OS
    rows = 5
    cols = 3
    left = Inches(0.8)
    top = Inches(2.2)
    width = Inches(11.7)
    height = Inches(4.5)
    
    table_shape = s10.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(4.25)
    table.columns[2].width = Inches(4.25)

    headers_text = ["Evaluation Dimension", "Traditional Manual Operation", "Nexus OS Resource Guardian"]
    for j, htext in enumerate(headers_text):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(241, 245, 249)
        p = cell.text_frame.paragraphs[0]
        p.text = htext
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK

    data_matrix = [
        ("Response Time", "Minutes to hours (after system freezes)", "Under 2 seconds (Autonomous)"),
        ("Problem Prevention", "Reactive (fixes after system crash)", "Proactive (10-min predictive horizon)"),
        ("System Safety", "High risk of manual operator error", "Deterministic immunity blacklist guardrails"),
        ("Governance & Audit", "Fragmented raw text logs", "Formatted PDF report with SHA-256 signature")
    ]

    for i, row_data in enumerate(data_matrix):
        for j, val in enumerate(row_data):
            cell = table.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(248, 250, 252)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(12)
            p.font.color.rgb = TEXT_DARK if j == 2 else TEXT_MUTED
            if j == 2:
                p.font.bold = True

    output_filename = "Nexus_OS_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully to {output_filename}")

if __name__ == "__main__":
    create_nexus_presentation()
